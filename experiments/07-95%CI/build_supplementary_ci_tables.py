import argparse
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


HERE = Path(__file__).resolve().parent
DEFAULT_MAIN_RESULTS = HERE / "outputs" / "validation" / "validation_metrics_95ci.csv"
DEFAULT_CRPA_RESULTS = (
    HERE / "outputs" / "crpa_foldwise_oof" / "crpa_foldwise_bootstrap_95ci.csv"
)
DEFAULT_OUTPUT = HERE / "outputs" / "supplementary_ci_tables.docx"

METRICS = ["Accuracy", "Precision", "Recall", "F1-score", "AUROC", "AUPRC"]
MODEL_ORDER_7 = [
    "XGBoost",
    "LightGBM",
    "CatBoost",
    "RandomForest",
    "LogisticRegression",
    "SVM",
    "KNN",
]
MODEL_ORDER_5 = MODEL_ORDER_7[:5]
TASK_LABELS = {
    "clinical_efficacy": "Clinical efficacy",
    "survival": "Survival",
    "polymyxin_resistance": "Polymyxin resistance",
}
TEMP_FEATURE_ORDER = {
    "clinical_efficacy": [
        "Full feature set",
        "Feature-set A",
        "Feature-set B",
        "Feature-set C",
        "Feature-set D",
    ],
    "survival": ["Full feature set", "Feature-set A", "Feature-set B"],
    "polymyxin_resistance": [
        "Full feature set",
        "Feature-set A",
        "Feature-set B",
        "Feature-set C",
        "Feature-set D",
        "Feature-set E",
    ],
}


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep_next = OxmlElement("w:keepNext")
    p_pr.append(keep_next)


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    keep_with_next(p)
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(9)


def format_metric(estimate, lower, upper):
    return f"{estimate:.3f}\n({lower:.3f}–{upper:.3f})"


def pivot_main(data):
    index = ["cohort", "task", "feature_mode", "candidate", "feature_set", "model"]
    wide = data.pivot_table(index=index, columns="metric", values=["estimate", "ci_lower", "ci_upper"]).reset_index()
    wide.columns = ["__".join([str(x) for x in col if str(x) != ""]) for col in wide.columns]
    return wide


def metric_value(row, metric):
    return format_metric(
        float(row[f"estimate__{metric}"]),
        float(row[f"ci_lower__{metric}"]),
        float(row[f"ci_upper__{metric}"]),
    )


def style_table(table, widths):
    table.style = "Table Grid"
    table.autofit = False
    header = table.rows[0]
    repeat_header(header)
    for j, cell in enumerate(header.cells):
        cell.width = widths[j]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "D9EAF7")
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            for r in p.runs:
                r.bold = True
                r.font.name = "Times New Roman"
                r.font.size = Pt(7)
    for row in table.rows[1:]:
        prevent_row_split(row)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for j, cell in enumerate(row.cells):
            cell.width = widths[j]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(7)


def add_note(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(7)


def add_performance_table(doc, records, include_task=False):
    headers = (["Task"] if include_task else []) + ["Feature set", "Model"] + [
        f"{m}\n(95% CI)" for m in METRICS
    ]
    table = doc.add_table(rows=1, cols=len(headers))
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for record in records:
        cells = table.add_row().cells
        values = ([record["task"]] if include_task else []) + [record["feature_set"], record["model"]] + [
            record[m] for m in METRICS
        ]
        for cell, value in zip(cells, values):
            cell.text = value
    widths = ([Cm(2.0)] if include_task else []) + [Cm(1.65), Cm(2.15)] + [Cm(2.25)] * 6
    style_table(table, widths)
    return table


def prepare_records(frame, task=None, crpa=False):
    records = []
    if crpa:
        candidate_labels = {
            "all_features": "Full feature set",
            "cand_04": "Feature-set A",
            "cand_25": "Feature-set B",
        }
        for candidate in ["all_features", "cand_04", "cand_25"]:
            for model in MODEL_ORDER_5:
                block = frame[(frame["candidate"] == candidate) & (frame["model"] == model)]
                if block.empty:
                    continue
                rec = {"feature_set": candidate_labels[candidate], "model": model}
                for metric in METRICS:
                    row = block[block["metric"] == metric].iloc[0]
                    rec[metric] = format_metric(row["estimate"], row["ci_lower"], row["ci_upper"])
                records.append(rec)
        return records

    if task is not None:
        frame = frame[frame["task"] == task]
    if task == "survival" and (frame["cohort"] == "external_mimic_crab").all():
        feature_labels = {
            "Feature-set A": "feature-setA",
            "Feature-set B": "feature-setB",
        }
        for label, raw in feature_labels.items():
            for model in MODEL_ORDER_7:
                block = frame[(frame["feature_set"] == raw) & (frame["model"] == model)]
                if block.empty:
                    continue
                rec = {"feature_set": label, "model": model}
                for metric in METRICS:
                    row = block[block["metric"] == metric].iloc[0]
                    rec[metric] = format_metric(row["estimate"], row["ci_lower"], row["ci_upper"])
                records.append(rec)
        return records

    candidate_to_label = dict(
        frame[["candidate", "feature_set"]]
        .drop_duplicates()
        .assign(
            label=lambda x: x["feature_set"].map(
                lambda value: "Full feature set"
                if value == "all_features"
                else value.replace("feature-set", "Feature-set ")
            )
        )[["candidate", "label"]]
        .itertuples(index=False, name=None)
    )
    desired = TEMP_FEATURE_ORDER[task]
    for label in desired:
        candidate = next(key for key, value in candidate_to_label.items() if value == label)
        for model in MODEL_ORDER_7:
            block = frame[(frame["candidate"] == candidate) & (frame["model"] == model)]
            if block.empty:
                continue
            rec = {"task": TASK_LABELS[task], "feature_set": label, "model": model}
            for metric in METRICS:
                row = block[block["metric"] == metric].iloc[0]
                rec[metric] = format_metric(row["estimate"], row["ci_lower"], row["ci_upper"])
            records.append(rec)
    return records


def parse_args():
    parser = argparse.ArgumentParser(
        description="Build the standalone supplementary 95% CI table document."
    )
    parser.add_argument(
        "--main-results", type=Path, default=DEFAULT_MAIN_RESULTS,
        help="Validation CI CSV produced by run_validation_ci.py.",
    )
    parser.add_argument(
        "--crpa-results", type=Path, default=DEFAULT_CRPA_RESULTS,
        help="CRPA foldwise CI CSV produced by run_crpa_ci_foldwise_oof.py.",
    )
    parser.add_argument(
        "--output", type=Path, default=DEFAULT_OUTPUT,
        help="Destination DOCX path.",
    )
    return parser.parse_args()


def main():
    args = parse_args()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    main_data = pd.read_csv(args.main_results)
    crpa_data = pd.read_csv(args.crpa_results)

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(1.0)
    section.bottom_margin = Cm(1.0)
    section.left_margin = Cm(1.0)
    section.right_margin = Cm(1.0)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(8)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Complete validation performance with 95% confidence intervals")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)

    crab = main_data[
        (main_data["cohort"] == "external_mimic_crab")
        & (main_data["task"] == "survival")
        & (main_data["metric"].isin(METRICS))
    ]
    add_caption(
        doc,
        "Supplementary Table 4. Complete model-specific performance for survival prediction in the external MIMIC-IV CRAB cohort with 95% confidence intervals.",
    )
    add_performance_table(doc, prepare_records(crab, task="survival"))
    add_note(
        doc,
        "Values are point estimates (95% CIs). CIs were obtained by patient-level non-parametric percentile bootstrap resampling of pooled LOOCV out-of-fold predictions (10,000 valid replicates). The classification threshold was fixed at 0.5.",
    )

    doc.add_page_break()
    temporal = main_data[
        (main_data["cohort"] == "temporal_external-1")
        & (main_data["task"].isin(TASK_LABELS))
        & (main_data["metric"].isin(METRICS))
    ]
    add_caption(
        doc,
        "Supplementary Table 5. Complete model-specific performance in the temporal validation cohort with 95% confidence intervals.",
    )
    temp_records = []
    for task in ["clinical_efficacy", "survival", "polymyxin_resistance"]:
        temp_records.extend(prepare_records(temporal, task=task))
    add_performance_table(doc, temp_records, include_task=True)
    add_note(
        doc,
        "Values are point estimates (95% CIs). CIs were obtained by patient-level non-parametric percentile bootstrap resampling of fixed temporal-cohort predictions (10,000 valid replicates). The classification threshold was fixed at 0.5.",
    )

    doc.add_page_break()
    add_caption(
        doc,
        "Supplementary Table 6. Complete model-specific performance for 28-day mortality prediction in the external MIMIC-IV CRPA cohort with 95% confidence intervals.",
    )
    add_performance_table(doc, prepare_records(crpa_data, crpa=True))
    add_note(
        doc,
        "Values are the original five-fold mean estimates (95% CIs). CIs were obtained by within-fold patient-level bootstrap resampling of fixed out-of-fold predictions, with each replicate summarized as the mean of the five fold-specific metrics (10,000 valid replicates). The classification threshold was fixed at 0.5.",
    )

    props = doc.core_properties
    props.title = "Supplementary Tables 4–6: complete validation performance with 95% CIs"
    props.subject = "Reviewer 2, Comment 2"
    doc.save(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
